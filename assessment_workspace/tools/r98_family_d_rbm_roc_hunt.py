"""HUNT script for R-98 Family D: RBM + ROC (Rules Service) rows."""
from __future__ import annotations
import hashlib, json
from pathlib import Path
from openpyxl import load_workbook

SRR_PATH = Path("/app/assessment_workspace/custody/owner_drops/20260920T1900Z/Master_Requirements_Registry_v5.xlsx")
EXPECTED_SHA = "00072c5829d3d89399cc5ca64661b405ac8a3cd0b43d0081f0d5f0bce109619d"

def main():
    sha = hashlib.sha256(SRR_PATH.read_bytes()).hexdigest()
    assert sha == EXPECTED_SHA
    wb = load_workbook(SRR_PATH, read_only=True, data_only=True)
    ws = wb["System Rules Registry"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    print("HEADERS:", headers)

    rbm_rows = []
    roc_rows = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        rid = row[0]
        if isinstance(rid, str):
            rid_up = rid.strip().upper()
            if rid_up.startswith("RBM-"):
                rbm_rows.append((row_num, rid.strip(), row))
            elif rid_up.startswith("ROC-"):
                roc_rows.append((row_num, rid.strip(), row))
    wb.close()

    print(f"\nRBM rows: {len(rbm_rows)}")
    print(f"ROC rows: {len(roc_rows)}")

    def dump(rows, name):
        for rn, rid, r in rows:
            stmt = r[1] if len(r) > 1 else None
            promise = r[3] if len(r) > 3 else None
            status = r[5] if len(r) > 5 else None
            domain = r[7] if len(r) > 7 else None
            print(f'{rid:8s} row={rn} status={status} domain={domain!r} promise={promise!r}')
            if stmt:
                print(f'    stmt: {str(stmt)[:130]}')

    print("\n=== RBM ROWS ===")
    dump(rbm_rows, "RBM")
    print("\n=== ROC ROWS ===")
    dump(roc_rows, "ROC")

    # Full dump
    out = {
        "wave": "R-98-CANDIDATE-CP-G5V5-FAMILY-D-RBM-ROC-HUNT",
        "family": "D",
        "authority_pins": {
            "v5_srr_workbook_sha256": sha,
            "v5_srr_workbook_path": "custody/owner_drops/20260920T1900Z/Master_Requirements_Registry_v5.xlsx",
        },
        "total_rbm": len(rbm_rows),
        "total_roc": len(roc_rows),
        "rbm_rows": [{"srr_row": rn, "rule_id": rid, "statement": r[1], "objective": r[2], "promise": r[3], "doc": r[4], "status": r[5], "enforcement_class": r[6], "domain": r[7], "source_file_line": r[8], "relationship": r[9]} for rn, rid, r in rbm_rows],
        "roc_rows": [{"srr_row": rn, "rule_id": rid, "statement": r[1], "objective": r[2], "promise": r[3], "doc": r[4], "status": r[5], "enforcement_class": r[6], "domain": r[7], "source_file_line": r[8], "relationship": r[9]} for rn, rid, r in roc_rows],
        "hunt_at_utc": "2026-09-21T07:00:00Z",
    }
    out_path = Path("/app/assessment_workspace/phase0/spec_hunt_receipts/20260921T070000Z_r98_family_d_rbm_roc_hunt_data.json")
    out_path.write_text(json.dumps(out, indent=2, default=str, ensure_ascii=False))
    print(f"\nHUNT data banked at: {out_path}")

if __name__ == "__main__":
    main()
