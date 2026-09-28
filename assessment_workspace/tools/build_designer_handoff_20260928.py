"""Publish only the explicit nonsecret designer brief/status, never credential files."""
import json
import shutil
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('/app')
SOURCE = ROOT / 'assessment_workspace/notes/designer_handoff_20260928'
PUBLIC = ROOT / 'frontend/public/briefs'


def word_document(markdown, target):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('173E48')
    lines, index = markdown.splitlines(), 0
    while index < len(lines):
        line = lines[index].strip()
        if line.startswith('|'):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells = [c.strip() for c in lines[index].strip().strip('|').split('|')]
                if not all(set(c) <= set('-: ') for c in cells):
                    table_lines.append(cells)
                index += 1
            table = doc.add_table(rows=0, cols=len(table_lines[0]))
            table.style = 'Light Shading Accent 1'
            for cells in table_lines:
                row = table.add_row()
                for cell, text in zip(row.cells, cells):
                    cell.text = text.replace('**', '').replace('`', '')
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:], 0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], 1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], 2)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:].replace('**', ''), 'List Bullet')
        elif line.startswith('> '):
            doc.add_paragraph(line[2:].replace('**', ''), 'Quote')
        elif line:
            paragraph = doc.add_paragraph()
            for i, part in enumerate(line.replace('`', '').split('**')):
                paragraph.add_run(part).bold = bool(i % 2)
        index += 1
    doc.core_properties.title = 'Akki — Product Experience Brief'
    doc.core_properties.subject = 'Outcome-led user stories, journeys and experience acceptance'
    doc.core_properties.author = 'Akki product team'
    doc.save(target)


def main():
    PUBLIC.mkdir(parents=True, exist_ok=True)
    mapping = {'AKKI_EXPERIENCE_BRIEF.md': 'Akki_Designer_Brief.md', 'BUILD_STATUS.md': 'Akki_Build_Status.md'}
    for src, dst in mapping.items():
        shutil.copyfile(SOURCE/src, PUBLIC/dst)
    word_document((SOURCE/'AKKI_EXPERIENCE_BRIEF.md').read_text(), PUBLIC/'Akki_Designer_Brief.docx')
    baseline = ROOT/'assessment_workspace/notes/full_requirements_gap_20260926/REQUIREMENT_GAPS_648.json'
    rows = json.loads(baseline.read_text())
    assert len(rows) == len({row['composite_id'] for row in rows}) == 648
    progress = {'as_of': datetime.now(timezone.utc).isoformat(),
                'completion_percent': None, 'accepted_requirement_numerator': None,
                'adopted_completion_denominator': None,
                'reason': 'Current requirement-level implementation/integration/acceptance reconciliation is incomplete.',
                'historical_assessment_date': '2026-09-26',
                'historical_registered_rows': len(rows), 'historical_families': len({r['doc'] for r in rows}),
                'historical_dispositions': dict(Counter(r['delivery_status'] for r in rows)),
                'historical_counts_are_current_completion': False,
                'last_verified_regression': {'backend_passed': 712, 'frontend_passed': 31, 'gated_skips': 2,
                                            'backend_report': '/app/test_reports/iteration_53.json',
                                            'frontend_unchanged_since_previous_checkpoint': True},
                'test_count_is_completion_percent': False,
                'all_backend_integrations_complete': False,
                'holistic_component_handoff_acceptance_run': False,
                'gpu_parked': True}
    (SOURCE/'PROGRESS_EVIDENCE.json').write_text(json.dumps(progress, indent=2)+'\n')
    print('Published standalone designer brief (Word/Markdown) and separate build-status document; no credentials exported.')


if __name__ == '__main__':
    main()