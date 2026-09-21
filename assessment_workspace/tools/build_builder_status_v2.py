#!/usr/bin/env python3
"""
Fill Builder Status column K on all 602 System Rules Registry rows.

v2 — content-based rule-ID scan, Doc-context disambiguation for canon anchors,
plus service-package fallback for PARTIAL.
"""
from __future__ import annotations
import os
import re
import json
import subprocess
import hashlib
from pathlib import Path
from openpyxl import load_workbook

SRC = "/app/audit/scratch/Master_Requirements_Registry_QA_v2.xlsx"
DST = "/app/audit/inbox/Master_Requirements_Registry_QA_v2_builder_status.xlsx"
REPO = "/app/assessment_workspace/repos/Extraction-Akki"
TESTS_DIR = f"{REPO}/backend/tests/invariants"
SERVICES_DIR = f"{REPO}/backend/services"
BASELINE_FAILS = f"{REPO}/backend/tests/invariants/baseline_failures.json"

def _git(*args: str, cwd: str = REPO) -> str:
    r = subprocess.run(("git",) + args, cwd=cwd, capture_output=True, text=True)
    return r.stdout.strip()

HEAD_SHA = _git("rev-parse", "HEAD")

# ---------------------------------------------------------------------------
# Load all invariant test file contents ONCE.
# ---------------------------------------------------------------------------
test_files: dict[str, str] = {}
for p in Path(TESTS_DIR).glob("test_*.py"):
    try:
        test_files[p.name] = p.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        pass

# baseline_failures = set of failing test qualified names
_baseline: set[str] = set()
_baseline_files: set[str] = set()
with open(BASELINE_FAILS) as f:
    _bd = json.load(f)
for line in _bd["baseline_failures"]:
    _baseline.add(line)
    fname = line.split("::")[0].split("/")[-1]
    _baseline_files.add(fname)

# services present
_all_services = set()
for entry in Path(SERVICES_DIR).iterdir():
    n = entry.name
    if n.startswith("__"): continue
    if n.endswith(".pyc"): continue
    _all_services.add(n.replace(".py", "").lower())

# ---------------------------------------------------------------------------
# Domain → service package hints
# ---------------------------------------------------------------------------
DOMAIN_TO_SVC = {
    "synisense": ["synisense", "canon_doc1_chokepoint", "chokepoint"],
    "solva": ["solva", "solva_depth", "solva_v2"],
    "northena": ["northena", "canon_doc5_northena"],
    "mtafiti": ["mtafiti"],
    "targeta": ["targeta"],
    "mkondo": ["mkondo", "handoff"],
    "watch": ["watch", "n_series"],
    "ledger": ["ledger_writer", "ledger_verifier", "ledger_emitter", "cost_ledger"],
    "checker": ["checker", "critic_pass"],
    "registry": ["registry", "resource_registry", "capability_registry"],
    "outer gate": ["outer_gate"],
    "cost": ["compute_estate", "cost_ledger", "compute_seam", "economics"],
    "training": ["training", "training_runner", "shield_train"],
    "wpw": ["wpw", "whitepapers"],
    "dap": ["dap", "dap_lifecycle", "canon_dap_lifecycle"],
    "entitlement": ["entitlement_service_v1"],
    "memory": ["memory"],
    "principal": ["principal_store"],
    "content": ["content_registry"],
    "rules": ["rules_service", "canon_rules_service"],
    "governance": ["governance", "governance_plane"],
    "agent containment": ["agent_containment"],
    "orchestration": ["orchestrator_resilience", "release_gate"],
    "audit-trail": ["audit_trail"],
    "response sanitizer": ["response_sanitizer"],
    "security": ["security_function_policies", "serving_floor", "secret_hygiene"],
    "compliance": ["compliance", "rights_quarantine"],
    "auth": ["auth", "ip_allowlist"],
    "documents": ["documents_registry"],
    "deck rendering": ["opportunity_briefs", "artifact_store"],
    "critic": ["critic_pass"],
    "doc 0": ["doc0_attributes"],
    "trust domain": ["trust_domain"],
    "declaration": ["declaration"],
    "enforce": ["enforce"],
    "connect": ["connect", "connect_estate"],
    "census": ["census_canon", "census_dimensions", "census_engine"],
    "audio": ["audio_plane_v1", "audio_poc"],
    "data": ["data", "data_source", "data_suite_v2"],
    "release": ["release_gate", "release_budget"],
    "credential": ["credential_plane", "credentials_paste", "kms"],
    "team": ["team"],
    "approvals": ["approvals"],
    "gateway": ["gateway"],
    "personal client": ["personal_client"],
    "answering": ["answering_service"],
    "chat": ["chat"],
    "wizard": ["wizard"],
    "master_admin": ["master_admin"],
    "layer a": ["layer_a"], "layer b": ["layer_b"], "layer c": ["layer_c"],
    "perception": ["perception"],
    "grey mirror": ["grey_mirror"],
    "reliable time": ["reliable_time_source"],
    "campaign": ["campaign_infra"],
    "observability": ["observability"],
    "canonical projection": ["canonical_projection"],
    "attribute stream bundle": ["attribute_stream_bundle"],
    "n series": ["n_series"],
    "business continuity": ["business_continuity"],
    "evidence": ["evidence_register"],
    "baseline line": ["baseline_line"],
    "rights quarantine": ["rights_quarantine"],
    "estate keys": ["estate_keys"],
    "secret hygiene": ["secret_hygiene"],
    "cif": ["cif", "chokepoint"],
    "chokepoint": ["chokepoint", "canon_doc1_chokepoint"],
}

def resolve_domain(domain: str) -> list[str]:
    if not domain:
        return []
    d = str(domain).strip().lower()
    for key, pkgs in DOMAIN_TO_SVC.items():
        if key in d or (len(key) > 5 and key in d):
            return [p for p in pkgs if p in _all_services]
    return []

# ---------------------------------------------------------------------------
# Build search patterns for a rule (verbatim substring hunt)
# ---------------------------------------------------------------------------

def id_variants(rule_id: str) -> list[str]:
    """Verbatim-substring variants of a rule ID as they might appear in test content."""
    r = str(rule_id).strip()
    variants = {r}                                              # e.g., "ENT-3"
    variants.add(r.replace("-", ""))                            # "ENT3"
    variants.add(r.replace("-", "_"))                           # "ENT_3"
    if "." in r:
        variants.add(r.replace(".", "_"))                       # "A5_1"
    return sorted(v for v in variants if len(v) >= 3)

def canon_anchor_variants(rule_id: str, doc_n: int) -> list[tuple[str, str]]:
    """For A5.x / B5.x / C5.x rules, return the (doc-scoped) match patterns.
    Returns list of (regex_pattern, human_label) — must MATCH BOTH doc and identifier.
    """
    r = str(rule_id).strip()
    pats: list[tuple[str, str]] = []
    if re.match(r"^[ABC]5\.[0-9]+$", r) or re.match(r"^[ABC][0-9]*\.[0-9]+$", r):
        # doc must appear in same file
        v_dot = r
        v_us = r.replace(".", "_")
        # doc mentioned as "Doc N" or "doc_N" or "docN"
        pats.append((rf"(?i)\bdoc\s*_?\s*{doc_n}\b", "doc"))
        pats.append((rf"\b{re.escape(v_dot)}\b|\b{re.escape(v_us)}\b", "id"))
    return pats

# ---------------------------------------------------------------------------
# Test-scan: for each rule, find test files that mention it
# ---------------------------------------------------------------------------

def find_evidence(rule_id: str, doc: str) -> tuple[list[str], list[str]]:
    """Return (green_test_files, red_test_files) for a rule.
    - green = test files where the identifier appears, filename NOT in baseline_failures
    - red = test files where the identifier appears, filename IS in baseline_failures
    """
    green: list[str] = []
    red: list[str] = []

    # Extract doc number
    doc_n: int | None = None
    if doc:
        m = re.search(r"Doc\s*(\d+)", str(doc), re.I)
        if m:
            doc_n = int(m.group(1))

    variants = id_variants(rule_id)
    is_canon = bool(re.match(r"^[ABC]5?\.?[0-9]+$", str(rule_id)))
    canon_pats = canon_anchor_variants(rule_id, doc_n) if doc_n and is_canon else []

    # We need at least one variant match in the file. For canon anchors, require BOTH
    # doc mention AND identifier mention in the SAME file.
    for fname, content in test_files.items():
        hit = False
        if canon_pats:
            # both patterns must match
            all_match = True
            for pat, _ in canon_pats:
                if not re.search(pat, content):
                    all_match = False
                    break
            hit = all_match
        else:
            # verbatim substring — but must be word-boundaried on both sides
            # to avoid ENT-3 matching ENT-30
            for v in variants:
                # match on \b or [_ .:\'"] boundary
                if re.search(r"(^|[^A-Za-z0-9_])" + re.escape(v) + r"($|[^A-Za-z0-9])", content):
                    hit = True
                    break
            # If no verbatim hit, try filename-token match (for prefix-only IDs)
            if not hit:
                fn_low = fname.lower()
                for v in variants:
                    if v.lower() in fn_low:
                        hit = True
                        break
        if hit:
            (red if fname in _baseline_files else green).append(fname)

    return sorted(set(green)), sorted(set(red))

# ---------------------------------------------------------------------------
# Commit citation cache
# ---------------------------------------------------------------------------
_cache: dict[str, str] = {}

def commit_for(rel_path: str) -> str:
    """Return commit that added rel_path, else latest touching commit, else HEAD."""
    if rel_path in _cache:
        return _cache[rel_path]
    r = subprocess.run(
        ["git", "log", "--diff-filter=A", "-1", "--format=%H", "--", rel_path],
        cwd=REPO, capture_output=True, text=True,
    )
    sha = r.stdout.strip()
    if not sha:
        r2 = subprocess.run(
            ["git", "log", "-1", "--format=%H", "--", rel_path],
            cwd=REPO, capture_output=True, text=True,
        )
        sha = r2.stdout.strip()
    if not sha:
        sha = HEAD_SHA
    _cache[rel_path] = sha
    return sha

def _service_evidence(domain: str) -> tuple[str, str] | None:
    for pkg in resolve_domain(domain):
        for cand in (f"backend/services/{pkg}/__init__.py",
                     f"backend/services/{pkg}/main.py",
                     f"backend/services/{pkg}.py"):
            if Path(f"{REPO}/{cand}").exists():
                return (commit_for(cand), cand)
        pkg_dir = f"backend/services/{pkg}"
        if Path(f"{REPO}/{pkg_dir}").is_dir():
            return (commit_for(pkg_dir), pkg_dir)
    return None

def classify(rule_id: str, doc: str, domain: str, statement: str) -> tuple[str, str]:
    green, red = find_evidence(rule_id, doc)
    if green:
        primary = green[0]
        rel = f"backend/tests/invariants/{primary}"
        sha = commit_for(rel)
        note = f" (+{len(green)-1} more)" if len(green) > 1 else ""
        return "BUILT", f"commit:{sha[:12]} · {rel}{note}"
    if red:
        primary = red[0]
        rel = f"backend/tests/invariants/{primary}"
        sha = commit_for(rel)
        return "PARTIAL", f"commit:{sha[:12]} · {rel} (in baseline_failures)"
    svc = _service_evidence(domain)
    if svc:
        sha, path = svc
        return "PARTIAL", f"commit:{sha[:12]} · {path} (service scaffolding, no invariant naming this rule)"
    return "NOT BUILT", ""

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    wb = load_workbook(SRC, data_only=False, keep_links=False)
    ws = wb["System Rules Registry"]
    if ws.cell(row=1, column=11).value != "Builder Status":
        raise SystemExit("column K header mismatch — abort")

    distribution = {"BUILT": 0, "PARTIAL": 0, "NOT BUILT": 0}
    by_prefix: dict[str, dict[str, int]] = {}

    for r in range(2, ws.max_row + 1):
        rule_id = ws.cell(row=r, column=1).value
        if not rule_id:
            continue
        doc = ws.cell(row=r, column=5).value or ""
        domain = ws.cell(row=r, column=8).value or ""
        statement = ws.cell(row=r, column=2).value or ""
        status, citation = classify(str(rule_id), str(doc), str(domain), str(statement))
        cell_val = status if not citation else f"{status} · {citation}"
        ws.cell(row=r, column=11, value=cell_val)
        distribution[status] += 1
        m = re.match(r"^([A-Za-z]+)", str(rule_id))
        pfx = m.group(1) if m else "?"
        by_prefix.setdefault(pfx, {"BUILT": 0, "PARTIAL": 0, "NOT BUILT": 0})[status] += 1

    Path(DST).parent.mkdir(parents=True, exist_ok=True)
    wb.save(DST)
    print("Saved:", DST)
    print("Distribution:", distribution)
    print("Total:", sum(distribution.values()))
    print()
    print("By prefix:")
    for p in sorted(by_prefix):
        d = by_prefix[p]
        print(f"  {p:6s}  BUILT={d['BUILT']:4d}  PARTIAL={d['PARTIAL']:4d}  NOT BUILT={d['NOT BUILT']:4d}  total={sum(d.values())}")

    summary = {
        "distribution": distribution,
        "by_prefix": by_prefix,
        "head_sha": HEAD_SHA,
        "workbook_src_sha256": hashlib.sha256(open(SRC, "rb").read()).hexdigest(),
        "workbook_dst_sha256": hashlib.sha256(open(DST, "rb").read()).hexdigest(),
        "baseline_failing_test_files": sorted(_baseline_files),
        "invariant_test_count": len(test_files),
    }
    with open("/app/audit/scratch/builder_status_summary.json", "w") as f:
        json.dump(summary, f, indent=2)
    print("\nSummary written to /app/audit/scratch/builder_status_summary.json")
    print("Filled workbook sha256:", summary["workbook_dst_sha256"])

if __name__ == "__main__":
    main()
