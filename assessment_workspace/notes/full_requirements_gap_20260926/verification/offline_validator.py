import csv
import hashlib
import json
import random
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/app")
BASE = ROOT / "assessment_workspace/notes/full_requirements_gap_20260926"
OUT = BASE / "verification"


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def norm(s):
    return " ".join(str(s or "").split())


def list_tokens(s: str):
    stop = {
        "the",
        "and",
        "for",
        "with",
        "from",
        "that",
        "this",
        "under",
        "into",
        "only",
        "not",
        "are",
        "all",
        "can",
        "must",
        "will",
    }
    toks = re.findall(r"[a-z0-9_\-]+", (s or "").lower())
    return [t for t in toks if len(t) >= 4 and t not in stop]


def load_master_xlsx(master_path: Path):
    wb = load_workbook(master_path, data_only=True, read_only=True)
    ws = wb["System Rules Registry"] if "System Rules Registry" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = 1
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row]]
    hmap = {h: i for i, h in enumerate(headers)}
    out = []
    for r in range(2, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        rid = row[hmap.get("Rule ID", -1)] if "Rule ID" in hmap else None
        doc = row[hmap.get("Doc", -1)] if "Doc" in hmap else None
        stmt = row[hmap.get("Statement", -1)] if "Statement" in hmap else None
        if rid is None and doc is None and stmt is None:
            continue
        out.append(
            {
                "workbook_row": r,
                "doc": norm(doc),
                "rule_id": norm(rid),
                "statement": norm(stmt),
                "composite_id": f"{norm(doc)}::{norm(rid)}",
            }
        )
    return out


def load_promises_xlsx(promise_path: Path):
    wb = load_workbook(promise_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[2]]
    header_row = 2
    if "Promise" not in headers:
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        header_row = 1
    hmap = {h: i for i, h in enumerate(headers)}

    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        pid = row[hmap.get("#", -1)] if "#" in hmap else None
        promise = row[hmap.get("Promise", -1)] if "Promise" in hmap else None
        bound = row[hmap.get("Bound / Ruling", -1)] if "Bound / Ruling" in hmap else None
        if pid is None and promise is None:
            continue
        pid_n = norm(pid)
        # Contract scope is assurance promises S-1..S-35; ignore product #N rows.
        if not re.match(r"^S-\d+$", pid_n):
            continue
        out.append(
            {
                "workbook_row": r,
                "promise_id": pid_n,
                "promise_text": norm(promise),
                "bound_ruling": norm(bound),
            }
        )
    return out


def main():
    review_contract = read_json(BASE / "REVIEW_CONTRACT.json")
    req_gaps = read_json(BASE / "REQUIREMENT_GAPS_648.json")
    prom_gaps = read_json(BASE / "PROMISE_GAPS_35.json")
    baseline_648 = read_json(BASE / "BASELINE_648.json")
    baseline_prom = read_json(BASE / "BASELINE_PROMISES_35.json")
    coverage = read_json(BASE / "COVERAGE_RECEIPT.json")
    supplemental = read_json(BASE / "SUPPLEMENTAL_RECONCILIATION.json")
    crosswalk = read_json(BASE / "PRODUCT_PROMISE_CROSSWALK.json")
    report_text = (BASE / "REPORT.md").read_text(encoding="utf-8")
    req_md = (BASE / "REQUIREMENTS_LINE_BY_LINE.md").read_text(encoding="utf-8")
    prom_md = (BASE / "PROMISES_LINE_BY_LINE.md").read_text(encoding="utf-8")

    master_rel = review_contract["baseline"]["registry"]
    promise_rel = review_contract["baseline"]["promises"]
    master_path = ROOT / master_rel
    promise_path = ROOT / promise_rel

    contract_hash_checks = {
        "master_exists": master_path.exists(),
        "promise_exists": promise_path.exists(),
        "master_sha256_actual": sha256_file(master_path) if master_path.exists() else None,
        "master_sha256_expected": review_contract["baseline"]["registry_sha256"],
        "promise_sha256_actual": sha256_file(promise_path) if promise_path.exists() else None,
        "promise_sha256_expected": review_contract["baseline"]["promises_sha256"],
    }
    contract_hash_checks["master_hash_match"] = (
        contract_hash_checks["master_sha256_actual"] == contract_hash_checks["master_sha256_expected"]
    )
    contract_hash_checks["promise_hash_match"] = (
        contract_hash_checks["promise_sha256_actual"] == contract_hash_checks["promise_sha256_expected"]
    )

    master_rows = load_master_xlsx(master_path)
    promise_rows = load_promises_xlsx(promise_path)

    req_by_composite = {r["composite_id"]: r for r in req_gaps}
    req_comp_list = [r["composite_id"] for r in req_gaps]
    req_comp_dupes = [k for k, c in Counter(req_comp_list).items() if c > 1]

    xlsx_comp_list = [r["composite_id"] for r in master_rows]
    xlsx_comp_dupes = [k for k, c in Counter(xlsx_comp_list).items() if c > 1]

    req_comp_set = set(req_comp_list)
    xlsx_comp_set = set(xlsx_comp_list)

    missing_in_req = sorted(xlsx_comp_set - req_comp_set)
    extra_in_req = sorted(req_comp_set - xlsx_comp_set)

    statement_row_mismatches = []
    for row in master_rows:
        cid = row["composite_id"]
        g = req_by_composite.get(cid)
        if not g:
            continue
        if norm(g.get("statement")) != row["statement"] or int(g.get("workbook_row", -1)) != int(row["workbook_row"]):
            statement_row_mismatches.append(
                {
                    "composite_id": cid,
                    "xlsx_statement": row["statement"],
                    "gap_statement": norm(g.get("statement")),
                    "xlsx_row": row["workbook_row"],
                    "gap_row": g.get("workbook_row"),
                }
            )

    promise_id_dupes = [k for k, c in Counter([p["promise_id"] for p in prom_gaps]).items() if c > 1]
    promise_by_id = {p["promise_id"]: p for p in prom_gaps}
    xlsx_promise_by_id = {p["promise_id"]: p for p in promise_rows}
    missing_promises = sorted(set(xlsx_promise_by_id) - set(promise_by_id))
    extra_promises = sorted(set(promise_by_id) - set(xlsx_promise_by_id))

    promise_literal_mismatches = []
    for pid, prow in xlsx_promise_by_id.items():
        g = promise_by_id.get(pid)
        if not g:
            continue
        if norm(g.get("original_promise_text")) != norm(prow.get("promise_text")):
            promise_literal_mismatches.append(
                {
                    "promise_id": pid,
                    "field": "promise_text",
                    "xlsx": prow.get("promise_text"),
                    "gap": g.get("original_promise_text"),
                }
            )
        if norm(g.get("bound_ruling")) != norm(prow.get("bound_ruling")):
            promise_literal_mismatches.append(
                {
                    "promise_id": pid,
                    "field": "bound_ruling",
                    "xlsx": prow.get("bound_ruling"),
                    "gap": g.get("bound_ruling"),
                }
            )

    req_md_cards = len(re.findall(r"^## REQ-\d+", req_md, flags=re.M))
    prom_md_cards = len(re.findall(r"^## S-\d+", prom_md, flags=re.M))

    expected_status_counts = {
        "PARTIAL": 475,
        "UNVERIFIED": 128,
        "SOURCE_UNRESOLVED": 24,
        "SPEC_DECISION_REQUIRED": 7,
        "EXPLICIT_NONCLAIM": 1,
        "EVIDENCED_SCOPED": 13,
    }
    actual_status_counts = Counter([r.get("delivery_status") for r in req_gaps])

    code_issues = []
    sha_cache = {}

    def file_sha_cached(p: Path):
        if p not in sha_cache:
            sha_cache[p] = sha256_file(p)
        return sha_cache[p]

    def check_code_ev(parent_id, ev, category):
        rel = ev.get("path") or ev.get("file")
        if not rel:
            code_issues.append({"id": parent_id, "category": category, "issue": "missing path/file"})
            return
        p = ROOT / rel
        if not p.exists():
            code_issues.append({"id": parent_id, "category": category, "issue": "missing file", "path": rel})
            return
        actual = file_sha_cached(p)
        current_expected = ev.get("current_file_sha256")
        if current_expected and current_expected != actual:
            code_issues.append(
                {
                    "id": parent_id,
                    "category": category,
                    "issue": "current_file_sha256 mismatch",
                    "path": rel,
                    "expected": current_expected,
                    "actual": actual,
                }
            )
        try:
            lines = p.read_text(encoding="utf-8").splitlines()
        except Exception:
            lines = p.read_text(errors="ignore").splitlines()
        eof = len(lines)
        st = int(ev.get("start_line", 1) or 1)
        en = int(ev.get("end_line", st) or st)
        rev_en = int(ev.get("reviewer_end_line", en) or en)
        if st < 1 or st > eof:
            code_issues.append({"id": parent_id, "category": category, "issue": "start out of range", "path": rel, "start": st, "eof": eof})
        norm_end = min(max(en, rev_en), eof)
        if norm_end < st:
            code_issues.append(
                {
                    "id": parent_id,
                    "category": category,
                    "issue": "end out of range",
                    "path": rel,
                    "start": st,
                    "end": en,
                    "reviewer_end": rev_en,
                    "eof": eof,
                }
            )

    for r in req_gaps:
        for ev in r.get("code_evidence", []) or []:
            check_code_ev(r.get("review_id"), ev, "requirement")
    for p in prom_gaps:
        for ev in p.get("existing_mechanism_file_line_hash", []) or []:
            check_code_ev(p.get("promise_id"), ev, "promise")

    sup_rows = supplemental.get("rows", [])
    sup_class_counts = Counter([r.get("classification") for r in sup_rows])
    sup_expected_map = {
        "REGISTERED_OBLIGATION_OVERLAP": 214,
        "RULING_OR_AUTHORITY_DEPENDENCY": 64,
        "EXPLICIT_LIMIT_OR_NONCLAIM": 34,
        "SOURCE_VERSION_DUPLICATE": 26,
        "HISTORICAL_DIAGNOSIS_NOT_CURRENT_VERDICT": 1,
        "SOURCE_CONTEXT_REQUIRES_TRACE": 191,
        "FOUNDATIONAL_CANON_TRACE_PENDING": 318,
    }

    sup_missing_text = [r.get("supplement_id") for r in sup_rows if not norm(r.get("text"))]
    sup_missing_sha = [r.get("supplement_id") for r in sup_rows if not norm(r.get("source_sha256"))]

    src_hash_issues = []
    src_sha_cache = {}
    for r in sup_rows:
        src = r.get("source")
        if not src:
            src_hash_issues.append({"supplement_id": r.get("supplement_id"), "issue": "missing source path"})
            continue
        p = ROOT / src
        if not p.exists():
            src_hash_issues.append({"supplement_id": r.get("supplement_id"), "issue": "source path missing", "source": src})
            continue
        if p not in src_sha_cache:
            src_sha_cache[p] = sha256_file(p)
        if r.get("source_sha256") and r.get("source_sha256") != src_sha_cache[p]:
            src_hash_issues.append(
                {
                    "supplement_id": r.get("supplement_id"),
                    "issue": "source hash mismatch",
                    "source": src,
                    "expected": r.get("source_sha256"),
                    "actual": src_sha_cache[p],
                }
            )

    # semantic specials
    s10 = promise_by_id.get("S-10", {}).get("current_status_scope_confidence_authority", {}).get("status")
    s11 = promise_by_id.get("S-11", {}).get("current_status_scope_confidence_authority", {}).get("status")

    rpo_rows = [
        {
            "review_id": r.get("review_id"),
            "statement": r.get("statement"),
            "delivery_status": r.get("delivery_status"),
        }
        for r in req_gaps
        if any(k in (r.get("statement") or "").lower() for k in ["rpo", "site-loss", "site loss"])
    ]

    ep_rows = [
        {
            "review_id": r.get("review_id"),
            "rule_id": r.get("rule_id"),
            "delivery_status": r.get("delivery_status"),
            "statement": r.get("statement"),
        }
        for r in req_gaps
        if re.match(r"^EP-?\d+$", str(r.get("rule_id", "")).strip())
    ]

    r32_report_present = "STATE/DECISIONS.md:26004" in report_text and "R32" in report_text

    generic_gap_counter = Counter([norm(r.get("exact_gap")) for r in req_gaps if norm(r.get("exact_gap"))])
    generic_closure_counter = Counter([norm(r.get("closure_demonstration")) for r in req_gaps if norm(r.get("closure_demonstration"))])
    repeated_gaps = [{"text": k, "count": v} for k, v in generic_gap_counter.items() if v >= 12][:10]
    repeated_closures = [{"text": k, "count": v} for k, v in generic_closure_counter.items() if v >= 12][:10]

    # legacy held/proposed not auto-missing
    authority_buckets = defaultdict(Counter)
    for r in req_gaps:
        authority_buckets[norm(r.get("registry_authority_status"))][r.get("delivery_status")] += 1

    # deterministic 33-sample semantic spot-check
    seed_hex = review_contract["baseline"]["registry_sha256"]
    rng = random.Random(int(seed_hex, 16))
    all_idx = list(range(len(req_gaps)))
    sample_idx = sorted(rng.sample(all_idx, 33))

    sample_reviews = []
    for i in sample_idx:
        row = req_gaps[i]
        rid = row.get("review_id")
        statement = row.get("statement", "")
        statement_tokens = set(list_tokens(statement))

        code_present = False
        keyword_hit = 0
        snippet_preview = ""
        ce = (row.get("code_evidence") or [])
        if ce:
            ev = ce[0]
            p = ROOT / ev.get("path", "")
            if p.exists():
                code_present = True
                try:
                    lines = p.read_text(encoding="utf-8").splitlines()
                except Exception:
                    lines = p.read_text(errors="ignore").splitlines()
                st = max(1, int(ev.get("start_line", 1) or 1))
                en = min(len(lines), max(int(ev.get("reviewer_end_line", ev.get("end_line", st)) or st), st))
                clip = "\n".join(lines[st - 1 : min(en, st + 39)])
                snippet_preview = clip[:260]
                code_tokens = set(list_tokens(clip + " " + (ev.get("observation") or "")))
                keyword_hit = len(statement_tokens.intersection(code_tokens))

        src_interp = " ".join([str(x.get("interpretation", "")) for x in (row.get("additional_source_evidence") or [])])
        src_hit = len(statement_tokens.intersection(set(list_tokens(src_interp))))

        repeated_closure = generic_closure_counter[norm(row.get("closure_demonstration"))] >= 20
        repeated_gap = generic_gap_counter[norm(row.get("exact_gap"))] >= 20

        verdict = "agree"
        reason = "Code/source anchors are present and semantically aligned at scoped level."
        if not code_present:
            verdict = "disagree"
            reason = "Referenced code anchor file missing; scoped claim cannot be verified."
        elif keyword_hit == 0 and src_hit == 0:
            verdict = "uncertain"
            reason = "Anchor exists but lexical/semantic linkage to statement is weak in sampled slice."
        elif repeated_closure or repeated_gap:
            verdict = "uncertain"
            reason = "Closure/gap text repetition suggests template reuse; needs row-specific semantic validation."

        sample_reviews.append(
            {
                "review_id": rid,
                "composite_id": row.get("composite_id"),
                "workbook_row": row.get("workbook_row"),
                "verdict": verdict,
                "reason": reason,
                "codeproof_scope_vs_broad_promise_test": "Scoped static citation can prove inspected shape only; not full end-to-end promise delivery.",
                "keyword_overlap_score": keyword_hit,
                "source_overlap_score": src_hit,
                "code_snippet_preview": snippet_preview,
            }
        )

    sample_counts = Counter([s["verdict"] for s in sample_reviews])

    report_compact = re.sub(r"\s+", "", report_text).lower()
    report_checks = {
        "mentions_509_trace_closure_needed": "509sourceblocksstillneedsemantictraceclosure" in report_compact,
        "mentions_24_source_cases_unresolved": "24original-sourcecasesremainunresolved" in report_compact,
        "says_not_fully_closed_canon": "wouldbefalsetolabeltheentirereferenced-canonauditfullyclosed" in report_compact,
        "no_platform_percentage_claim": "nojustifiedwhole-platformdeliverypercentage" in report_compact,
        "mentions_old_146_tests_narrow_scope": "existing146passingtestsbelongtoanarrowerpreviousdap/censusslice" in report_compact,
        "mentions_mocked": "mocked" in report_compact,
        "contains_percent_symbol": "%" in report_text,
        "contains_unsupported_platform_percent_phrase": "unsupportedplatform" in report_compact,
    }

    crosswalk_unmapped_ok = sorted(crosswalk.get("numbers_1_to_32_without_direct_hash_reference", [])) == ["P-05", "P-28", "P-29"]

    mechanical_pass = all(
        [
            contract_hash_checks["master_hash_match"],
            contract_hash_checks["promise_hash_match"],
            len(req_gaps) == 648,
            len(prom_gaps) == 35,
            len(missing_in_req) == 0,
            len(extra_in_req) == 0,
            len(req_comp_dupes) == 0,
            len(xlsx_comp_dupes) == 0,
            len(statement_row_mismatches) == 0,
            len(missing_promises) == 0,
            len(extra_promises) == 0,
            len(promise_id_dupes) == 0,
            len(promise_literal_mismatches) == 0,
            req_md_cards == 648,
            prom_md_cards == 35,
            dict(actual_status_counts) == expected_status_counts,
            coverage.get("status_counts") == expected_status_counts,
            len(code_issues) == 0,
            len(sup_rows) == 848,
            dict(sup_class_counts) == sup_expected_map,
            len(sup_missing_text) == 0,
            len(sup_missing_sha) == 0,
            len(src_hash_issues) == 0,
            report_checks["mentions_509_trace_closure_needed"],
            report_checks["mentions_24_source_cases_unresolved"],
            report_checks["says_not_fully_closed_canon"],
            crosswalk_unmapped_ok,
            s10 == "PARTIAL",
            s11 == "PARTIAL",
            r32_report_present,
        ]
    )

    receipt = {
        "task": "final_read_only_assessment_verification",
        "mechanical_pass": mechanical_pass,
        "semantic_status": "incomplete_audit_with_preserved_uncertainty",
        "contract_hash_checks": contract_hash_checks,
        "xlsx_alignment": {
            "master_rows_count": len(master_rows),
            "requirement_rows_count": len(req_gaps),
            "missing_in_requirements": missing_in_req,
            "extra_in_requirements": extra_in_req,
            "duplicate_composite_ids_in_requirements": req_comp_dupes,
            "duplicate_composite_ids_in_xlsx": xlsx_comp_dupes,
            "statement_or_workbook_row_mismatches": statement_row_mismatches,
            "promise_rows_count": len(promise_rows),
            "promise_gap_rows_count": len(prom_gaps),
            "missing_promises": missing_promises,
            "extra_promises": extra_promises,
            "duplicate_promise_ids": promise_id_dupes,
            "promise_literal_mismatches": promise_literal_mismatches,
        },
        "rendered_cards": {"requirements_line_by_line": req_md_cards, "promises_line_by_line": prom_md_cards},
        "status_counts": {
            "expected": expected_status_counts,
            "actual_from_requirements": dict(actual_status_counts),
            "coverage_receipt": coverage.get("status_counts"),
            "note": "13 EVIDENCED_SCOPED rows are scoped evidence only, not full delivered promises.",
        },
        "code_evidence_integrity": {
            "issues_count": len(code_issues),
            "issues_sample": code_issues[:40],
        },
        "supplemental_848_reconciliation": {
            "rows": len(sup_rows),
            "classification_counts_actual": dict(sup_class_counts),
            "classification_counts_expected": sup_expected_map,
            "missing_text_count": len(sup_missing_text),
            "missing_source_sha_count": len(sup_missing_sha),
            "source_hash_path_issues_count": len(src_hash_issues),
            "source_hash_path_issues_sample": src_hash_issues[:30],
            "509_trace_pending_confirmed": sup_class_counts.get("SOURCE_CONTEXT_REQUIRES_TRACE", 0)
            + sup_class_counts.get("FOUNDATIONAL_CANON_TRACE_PENDING", 0)
            == 509,
        },
        "crosswalk_checks": {
            "unmapped_product_promises": crosswalk.get("numbers_1_to_32_without_direct_hash_reference", []),
            "expected_unmapped": ["P-05", "P-28", "P-29"],
            "pass": crosswalk_unmapped_ok,
        },
        "semantic_policy_checks": {
            "s10_status": s10,
            "s11_status": s11,
            "bcr_rpo_site_loss_rows": rpo_rows,
            "ep_rows": ep_rows,
            "r32_state_decision_reference_present": r32_report_present,
            "report_language_checks": report_checks,
        },
        "generic_repetition_scan": {
            "top_repeated_exact_gap_entries": repeated_gaps,
            "top_repeated_closure_entries": repeated_closures,
        },
        "authority_status_distribution": {k: dict(v) for k, v in authority_buckets.items()},
        "sample_33": {
            "seed": seed_hex,
            "sample_size": 33,
            "verdict_counts": dict(sample_counts),
            "rows": sample_reviews,
        },
        "conclusion": {
            "mechanical_verification": "pass" if mechanical_pass else "fail",
            "semantic_uncertainty": "preserved",
            "incomplete_overall_canon_audit": True,
            "notes": [
                "This verifies accounting/integrity and scoped semantic sampling only.",
                "It does not claim full all-Canon semantic closure or runtime delivery proof.",
                "509 source blocks remain analyst trace closure work; 24 source-unresolved primary cases remain.",
            ],
        },
    }

    OUT.mkdir(parents=True, exist_ok=True)
    out_path = OUT / "final_verification_receipt.json"
    out_path.write_text(json.dumps(receipt, indent=2, ensure_ascii=False), encoding="utf-8")

    # lightweight machine-readable pass/fail summary for downstream agents
    summary_path = OUT / "final_verification_summary.json"
    summary = {
        "mechanical_pass": mechanical_pass,
        "sample_verdict_counts": dict(sample_counts),
        "status_counts": dict(actual_status_counts),
        "supplemental_rows": len(sup_rows),
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(str(out_path))
    print(str(summary_path))


if __name__ == "__main__":
    main()
