"""Read-only second owner drop ingestion. Preserve prior evidence before revision."""
import hashlib
import json
import shutil
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPORT = HERE.parent
ROOT = REPORT.parents[2]
CUSTODY = ROOT / "assessment_workspace/custody/owner_drops/enhancement_specs_18_20260926"
PRIOR = ROOT / "assessment_workspace/custody/owner_drops/enhancement_specs_20260926"


def dump(name, data):
    (HERE / name).write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str) + "\n")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    out = {s.title: [{"row": i, "values": list(v)} for i, v in enumerate(s.iter_rows(values_only=True), 1) if any(x is not None for x in v)] for s in wb}
    wb.close()
    return out


def records(rows):
    header = rows[0]["values"]
    return [{**dict(zip(header, r["values"])), "_row": r["row"]} for r in rows[1:] if r["values"][0]]


def delta(before, after, key):
    b, a = {key(r): r for r in before}, {key(r): r for r in after}
    return {
        "added": [a[k] for k in sorted(set(a)-set(b))],
        "removed": [b[k] for k in sorted(set(b)-set(a))],
        "changed": [{"key": k, "fields": {col: {"before": b[k].get(col), "after": a[k].get(col)} for col in set(a[k])|set(b[k]) if col != "_row" and a[k].get(col) != b[k].get(col)}}
                    for k in sorted(set(a)&set(b)) if any(col != "_row" and a[k].get(col) != b[k].get(col) for col in set(a[k])|set(b[k]))],
    }


for original, snapshot in [("INTEGRATION_REVIEW.md", "INTEGRATION_REVIEW_FILES17_SNAPSHOT.md"), ("EVIDENCE_MANIFEST.json", "EVIDENCE_MANIFEST_FILES17_SNAPSHOT.json")]:
    if not (HERE / snapshot).exists():
        shutil.copy2(REPORT / original, HERE / snapshot)

docs = []
for path in sorted(CUSTODY.iterdir()):
    if path.suffix not in {".docx", ".xlsx"}:
        continue
    info = {"name": path.name, "bytes": path.stat().st_size, "sha256": sha(path), "outcome": "processed"}
    if path.suffix == ".xlsx":
        data = workbook(path)
        dump(path.stem + ".json", data)
        info["nonempty_rows_per_sheet"] = {k: len(v) for k, v in data.items()}
    else:
        document = Document(path)
        blocks, lines = [], []
        for index, child in enumerate(document.element.body.iterchildren(), 1):
            if child.tag == qn("w:p"):
                p = Paragraph(child, document)
                if p.text.strip():
                    blocks.append({"block": index, "kind": "paragraph", "text": p.text})
                    lines.append(f"[B{index}] {p.text}")
            elif child.tag == qn("w:tbl"):
                rows = [[c.text for c in row.cells] for row in Table(child, document).rows]
                blocks.append({"block": index, "kind": "table", "rows": rows})
                lines.extend(f"[B{index}R{i}] " + " | ".join(r) for i, r in enumerate(rows, 1))
        dump(path.stem + ".json", blocks)
        (HERE / (path.stem + ".md")).write_text("\n\n".join(lines) + "\n")
        info["nonempty_blocks"] = len(blocks)
        info["tracked_insertions"] = len(document.element.xpath(".//w:ins"))
        info["tracked_deletions"] = len(document.element.xpath(".//w:del"))
    docs.append(info)

old = workbook(PRIOR / "Master_Requirements_Registry.xlsx")
new = workbook(CUSTODY / "Master_Requirements_Registry.xlsx")
r, q = records(new["System Rules Registry"]), records(new["QA Verification"])
key = lambda row: (row["Doc"], row["Rule ID"])
keys, qkeys = {key(x) for x in r}, {key(x) for x in q}
diff = delta(records(old["System Rules Registry"]), r, key)
security = workbook(CUSTODY / "Security_Assurance_Promise_Set.xlsx")
old_sec = workbook(PRIOR / "Security_Assurance_Promise_Set.xlsx")
sdiff = delta(records(old_sec["Security Assurance Set"]), records(security["Security Assurance Set"]), lambda row: row["#"])
dump("DOCUMENT_MANIFEST.json", {"recorded_at": datetime.now(timezone.utc).isoformat(), "source_document_date": "2026-09-27 as printed; custody date is separate", "archive_sha256": sha(CUSTODY / "files_18.zip"), "documents": docs, "previous_report_snapshot_sha256": sha(HERE / "INTEGRATION_REVIEW_FILES17_SNAPSHOT.md")})
dump("DELTA_FROM_FILES17.json", {"registry_rows": len(r), "registry_unique_keys": len(keys), "qa_rows": len(q), "qa_unique_keys": len(qkeys), "missing_qa": sorted(keys-qkeys), "extra_qa": sorted(qkeys-keys), "registry_delta": diff, "security_delta": sdiff})
print(json.dumps({"documents": len(docs), "registry_rows": len(r), "qa_rows": len(q), "missing_qa": len(keys-qkeys), "new_rule_ids": [x["Rule ID"] for x in diff["added"]], "changed_old_rules": len(diff["changed"]), "removed_old_rules": len(diff["removed"]), "new_promises": [x["#"] for x in sdiff["added"]], "changed_old_promises": len(sdiff["changed"])}, indent=2))