"""Iteration 80 — Akki OS owner order 2026-10-02 verification.

Covers five new deliverables (read-only + ONE memory message):
  1. Board admission floor for lwg v3
  2. Report pack (studio commission com-c51e900a05a4)
  3. Static downloads (unauthenticated)
  4. Durable job memory (one live LLM call)
  5. Estate-language status (read-only)

Guardrails from problem statement are strictly observed:
  * No POST to /training/lab/estate-language, /text-estate, /runs, /board/new,
    /board/{run}/actions, or /api/studio/producers/*.
  * Exactly ONE POST to /board/reinforcement-lwg-v3/messages.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import time

import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    # fall back to reading frontend/.env
    with open("/app/frontend/.env") as fh:
        for line in fh:
            if line.startswith("REACT_APP_BACKEND_URL="):
                BASE_URL = line.split("=", 1)[1].strip().rstrip("/")
                break
assert BASE_URL, "REACT_APP_BACKEND_URL not set"

EMAIL = "demo.operator@demo.rms.example.com"
PASSWORD = "demo-operator-pw"
COMMISSION_ID = "com-c51e900a05a4"


@pytest.fixture(scope="session")
def token():
    r = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": EMAIL, "password": PASSWORD},
        timeout=30,
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert "access_token" in data, data
    return data["access_token"]


@pytest.fixture(scope="session")
def auth(token):
    return {"Authorization": f"Bearer {token}"}


# --- AUTH -------------------------------------------------------------------
def test_login_returns_access_token(token):
    assert isinstance(token, str) and len(token) > 20


# --- BOARD ------------------------------------------------------------------
def test_board_lwg_v3_in_service(auth):
    r = requests.get(f"{BASE_URL}/api/training/lab/board", headers=auth, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    runs = data.get("runs") or []
    by_id = {run.get("run_id") or run.get("id"): run for run in runs}
    assert "reinforcement-lwg-v3" in by_id, list(by_id.keys())
    v3 = by_id["reinforcement-lwg-v3"]
    assert v3.get("status") == "In service", v3.get("status")
    meta_blob = json.dumps(v3).lower()
    assert "in service for luhya (wanga)" in meta_blob, "missing meta phrase"
    # 75% floor language
    assert "75" in meta_blob, "75% floor not present in run blob"
    assert "lwg" in meta_blob
    # earlier iterations remain refused
    for prior in ("reinforcement-lwg", "reinforcement-lwg-v2"):
        assert prior in by_id, prior
        assert by_id[prior].get("status") == "Refused", (prior, by_id[prior].get("status"))


# --- COMMISSION LISTING + DETAIL -------------------------------------------
def test_commissions_list_contains_estate_pack(auth):
    r = requests.get(f"{BASE_URL}/api/studio/commissions", headers=auth, timeout=30)
    assert r.status_code == 200, r.text
    body = r.json()
    coms = body.get("commissions") if isinstance(body, dict) else body
    assert coms is not None
    found = next((c for c in coms if c.get("commission_id") == COMMISSION_ID), None)
    assert found is not None, f"{COMMISSION_ID} not in list"
    release = found.get("release") or {}
    assert release.get("state") == "clear", release
    assert found.get("current_content_version") == 3, found.get("current_content_version")


@pytest.fixture(scope="session")
def commission_detail(auth):
    r = requests.get(f"{BASE_URL}/api/studio/commissions/{COMMISSION_ID}", headers=auth, timeout=30)
    assert r.status_code == 200, r.text
    body = r.json()
    com = body.get("commission") if isinstance(body, dict) and "commission" in body else body
    return com


def test_commission_detail_frontier_pipeline(commission_detail):
    com = commission_detail
    authoring = com.get("authoring") or {}
    assert authoring.get("pipeline") == "frontier", authoring
    assert (authoring.get("planner") or {}).get("model") == "gpt-6-astra", authoring.get("planner")
    assert authoring.get("model") == "claude-sonnet-5-5", authoring.get("model")
    review = authoring.get("review") or {}
    assert review.get("model") == "claude-fable-5-1", review
    assert review.get("verdict") == "accept", review


def test_commission_last_build_released(commission_detail):
    builds = commission_detail.get("builds") or []
    assert builds, "no builds"
    last = builds[-1]
    assert last.get("status") == "released", last.get("status")
    prove = last.get("prove") or {}
    assert prove.get("byte_identical") is True, prove
    qa = last.get("qa") or {}
    assert qa.get("status") == "PASS", qa
    outputs = last.get("outputs") or []
    # outputs may be list of dicts with 'name' or strings
    names = []
    for o in outputs:
        if isinstance(o, dict):
            names.append(o.get("name") or o.get("filename") or "")
        else:
            names.append(str(o))
    joined = " ".join(names).lower()
    for ext in ("docx", "pptx", "xlsx", "html"):
        assert ext in joined, f"{ext} missing in outputs: {names}"


# --- ARTIFACT DOWNLOADS -----------------------------------------------------
EXPECTED_HASHES = {
    "docx": ("dc8ee18fc5c83b089c876daad735e7ae1cd2bc3a334d8b45f206646b5bb6ebd2", 373792),
    "pptx": ("128fb5e6273ed700a4c8a758268d7b1955fa1839d7d5534efce7da04280768ca", 189208),
    "xlsx": ("81bc117063576c2355ae2eff87cec4bfc3a1b44bdd017ebfeb5cdb09ff043c87", 758818),
}


def _artifact_names(build):
    outputs = build.get("outputs") or []
    names = []
    for o in outputs:
        if isinstance(o, dict):
            n = o.get("name") or o.get("filename")
            if n:
                names.append(n)
        else:
            names.append(str(o))
    return names


@pytest.fixture(scope="session")
def artifact_bytes(auth, commission_detail):
    builds = commission_detail.get("builds") or []
    last = builds[-1]
    build_id = last.get("build_id") or last.get("id")
    names = _artifact_names(last)
    picks = {}
    for ext in ("docx", "pptx", "xlsx"):
        match = next((n for n in names if n.lower().endswith(f".{ext}")), None)
        assert match, f"no {ext} output among {names}"
        r = requests.get(
            f"{BASE_URL}/api/studio/commissions/{COMMISSION_ID}/builds/{build_id}/artifact/{match}",
            headers=auth,
            timeout=60,
        )
        assert r.status_code == 200, (ext, r.status_code, r.text[:200])
        assert len(r.content) > 0
        picks[ext] = (match, r.content, r.headers.get("content-type", ""))
    return picks


def test_docx_download_integrity(artifact_bytes):
    name, content, ctype = artifact_bytes["docx"]
    assert "word" in ctype or "officedocument" in ctype, ctype
    h = hashlib.sha256(content).hexdigest()
    assert h == EXPECTED_HASHES["docx"][0], h
    assert len(content) == EXPECTED_HASHES["docx"][1]


def test_pptx_download_integrity(artifact_bytes):
    name, content, ctype = artifact_bytes["pptx"]
    h = hashlib.sha256(content).hexdigest()
    assert h == EXPECTED_HASHES["pptx"][0], h
    assert len(content) == EXPECTED_HASHES["pptx"][1]


def test_xlsx_download_integrity(artifact_bytes):
    name, content, ctype = artifact_bytes["xlsx"]
    h = hashlib.sha256(content).hexdigest()
    assert h == EXPECTED_HASHES["xlsx"][0], h
    assert len(content) == EXPECTED_HASHES["xlsx"][1]


def test_docx_opens_and_has_expected_content(artifact_bytes):
    from docx import Document  # python-docx
    _, content, _ = artifact_bytes["docx"]
    doc = Document(io.BytesIO(content))
    headings = [p.text for p in doc.paragraphs]
    assert any("Plain English: the terms used in this pack" in t for t in headings), \
        "glossary heading missing"
    table_blob = []
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                table_blob.append(cell.text)
    joined = "\n".join(table_blob)
    assert "Reclaimable" in joined
    assert "Wanga" in joined


def test_pptx_opens_with_53_slides(artifact_bytes):
    from pptx import Presentation
    _, content, _ = artifact_bytes["pptx"]
    pres = Presentation(io.BytesIO(content))
    assert len(pres.slides) == 53, len(pres.slides)


def test_xlsx_opens_with_expected_sheets(artifact_bytes):
    import openpyxl
    _, content, _ = artifact_bytes["xlsx"]
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    expected = {"Plain English", "Decisions", "Audio by language", "News scripts by station", "Audio file evidence"}
    missing = expected - set(wb.sheetnames)
    assert not missing, f"missing sheets: {missing} (have {wb.sheetnames})"
    audio_sheet = wb["Audio file evidence"]
    rows = list(audio_sheet.iter_rows(values_only=True))
    assert 3900 <= len(rows) <= 4000, len(rows)
    # locate header row that contains 'Path in staged corpus' in column B
    header_idx = None
    for i, r in enumerate(rows[:10]):
        if len(r) > 1 and r[1] and "path in staged corpus" in str(r[1]).lower():
            header_idx = i
            break
    assert header_idx is not None, [rows[i] for i in range(min(5, len(rows)))]
    empty_b = sum(1 for r in rows[header_idx + 1:] if not (r[1] if len(r) > 1 else None))
    assert empty_b == 0, f"{empty_b} empty col-B rows"


# --- STATIC DOWNLOADS -------------------------------------------------------
STATIC = {
    "RMS_Estate_and_Language_Review.docx": EXPECTED_HASHES["docx"],
    "RMS_Estate_and_Language_Review.pptx": EXPECTED_HASHES["pptx"],
    "RMS_Estate_and_Language_Review.xlsx": EXPECTED_HASHES["xlsx"],
}


@pytest.mark.parametrize("fname", list(STATIC.keys()))
def test_static_download(fname):
    sess = requests.Session()
    # unauthenticated
    r = sess.get(f"{BASE_URL}/downloads/{fname}", timeout=60)
    assert r.status_code == 200, (fname, r.status_code)
    sha, size = STATIC[fname]
    assert len(r.content) == size, (fname, len(r.content))
    assert hashlib.sha256(r.content).hexdigest() == sha, fname


# --- ESTATE LANGUAGE STATUS -------------------------------------------------
def test_estate_language_status(auth):
    r = requests.get(f"{BASE_URL}/api/training/lab/estate-language", headers=auth, timeout=30)
    assert r.status_code == 200, r.text
    body = r.json()
    audio = body.get("audio") or {}
    assert audio.get("status") == "completed", audio.get("status")
    result = audio.get("result") or {}
    assert result.get("devices") == 2, result.get("devices")
    assert result.get("read") == 3958, result.get("read")
    compute = result.get("compute") or {}
    gpu_hours = compute.get("gpu_hours_total")
    assert gpu_hours is not None and 0.9 < float(gpu_hours) < 1.2, gpu_hours
    text_estate = body.get("text_estate") or {}
    assert text_estate.get("status") == "completed", text_estate.get("status")
    text = body.get("text") or {}
    assert text.get("status") in ("running", "completed"), text.get("status")


# --- JOB MEMORY (ONE live LLM call) -----------------------------------------
def test_job_memory_remembers_v1_v2(auth):
    r = requests.post(
        f"{BASE_URL}/api/training/lab/board/reinforcement-lwg-v3/messages",
        headers=auth,
        json={"text": "From your memory of this job only: what learning rate did v1 and v2 use and how did v2 end?"},
        timeout=180,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    msgs = body.get("messages") or body.get("thread") or []
    # pick the last akki message
    akki = [m for m in msgs if (m.get("role") == "akki" or m.get("author") == "akki")]
    assert akki, body
    last = akki[-1]
    assert last.get("model") == "gpt-6-astra", last.get("model")
    assert last.get("status") != "unavailable", last.get("status")
    text = (last.get("text") or last.get("body") or "").lower()
    # v2 ended by abstaining / stopping early
    assert ("v2" in text) or ("second" in text)
    # mention stop/abstain
    assert re.search(r"abstain|stop|halt|refus|short", text), text[:400]
