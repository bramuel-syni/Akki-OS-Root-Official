"""In-repo stand-in for the `xlsx` skill's recalc.py, reached through DOC_STUDIO_TOOLS_DIR (doc-studio v1.1, owner condition A).

Recalculates a workbook with headless LibreOffice (the environment of record, version pinned in the build manifest),
writes the recalculated bytes back in place, then scans every cached value for Excel error strings.
Prints one JSON object; `status` is `success` only when LibreOffice ran and no formula error remains."""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


def main(path: str) -> int:
    src = Path(path).resolve()
    report = {"file": str(src), "status": "error", "errors": [], "libreoffice": None}
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        report["reason"] = "LIBREOFFICE_UNAVAILABLE: no soffice binary on this host; formulas were not recalculated"
        print(json.dumps(report)); return 1
    try:
        report["libreoffice"] = subprocess.run([soffice, "--version"], capture_output=True, text=True, timeout=60).stdout.strip()
    except (OSError, subprocess.SubprocessError) as exc:
        report["reason"] = f"LIBREOFFICE_VERSION_UNREADABLE: {type(exc).__name__}"
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(src)]
        try:
            run = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        except subprocess.TimeoutExpired:
            report["reason"] = "LIBREOFFICE_TIMEOUT"; print(json.dumps(report)); return 1
        produced = Path(tmp) / src.name
        if run.returncode != 0 or not produced.exists():
            report["reason"] = f"LIBREOFFICE_CONVERT_FAILED rc={run.returncode}: {(run.stderr or run.stdout)[-300:]}"
            print(json.dumps(report)); return 1
        shutil.copyfile(produced, src)
    from openpyxl import load_workbook
    wb = load_workbook(src, data_only=True)
    formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    report["errors"].append({"sheet": ws.title, "cell": cell.coordinate, "value": cell.value})
    for ws in load_workbook(src, data_only=False).worksheets:
        for row in ws.iter_rows():
            formulas += sum(1 for c in row if isinstance(c.value, str) and c.value.startswith("="))
    report.update(formulas=formulas, status="success" if not report["errors"] else "errors_found")
    print(json.dumps(report))
    return 0 if not report["errors"] else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
