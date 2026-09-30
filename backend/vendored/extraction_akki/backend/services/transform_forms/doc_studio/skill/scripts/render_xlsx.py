"""Render a resolved workbook plan to .xlsx with openpyxl.

Inputs are the only hardcoded numbers; every derived cell is a formula that
references them, so the model recalculates when an input changes. The
Checks sheet cross-checks formula results against the content model's
calculations so a wrong reference cannot ship silently.
"""
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from studio import Content, load  # noqa: E402

BLUE = Font(name="Arial", color="0000FF")
BLACK = Font(name="Arial", color="000000")
BOLD = Font(name="Arial", bold=True)
HEAD = Font(name="Arial", bold=True, color="FFFFFF")
KEY = PatternFill("solid", fgColor="FFFF00")


def render(plan, out, content_path):
    content = Content(load(content_path))
    wb = Workbook()
    wb.remove(wb.active)
    pal = plan["look"]["palette"]
    headfill = PatternFill("solid", fgColor=pal["structure"])
    input_cells = {}
    options_layout = {"sheet": "Options", "options": ["O1", "O2", "O3"], "derived_row": 7}

    def header(ws, row, labels):
        for c, l in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=c, value=l)
            cell.font = HEAD
            cell.fill = headfill
            cell.alignment = Alignment(vertical="center")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    for sh in plan["sheets"]:
        ws = wb.create_sheet(sh["name"])
        if sh["kind"] == "cover":
            ws["A1"] = plan["title"]
            ws["A1"].font = Font(name="Arial", size=16, bold=True)
            ws["A2"] = f'{plan["organisation"]}  ·  {plan["spec_id"]}  ·  direction {plan["direction"]}'
            ws["A4"] = "Legend"
            ws["A4"].font = BOLD
            ws["A5"] = "Blue text"
            ws["A5"].font = BLUE
            ws["B5"] = "hardcoded input, receipt alongside; edit only these"
            ws["A6"] = "Black text"
            ws["B6"] = "formula; do not overwrite"
            ws["A7"] = "Yellow fill"
            ws["A7"].fill = KEY
            ws["B7"] = "key assumption the reader should confirm"
            ws["A9"] = "Sheets: Inputs (facts) → Options (model) → Series (exhibit data) → Checks (cross-checks against the content model)"
            widths(ws, [18, 90])
        elif sh["kind"] == "inputs":
            header(ws, 1, ["ID", "Label", "Value", "Unit", "Class", "Receipt", "Resolves to"])
            for r, f in enumerate(sh["facts"], start=2):
                ws.cell(row=r, column=1, value=f["id"])
                ws.cell(row=r, column=2, value=f["label"])
                c = ws.cell(row=r, column=3, value=f["value"])
                c.font = BLUE
                if f["unit"] == "pct":
                    c.number_format = "0.0%"
                elif f["unit"] in ("GBPm", "m"):
                    c.number_format = "#,##0.0"
                elif f["unit"] == "GBP":
                    c.number_format = "#,##0.00"
                elif isinstance(f["value"], float):
                    c.number_format = "#,##0.0"  # v1.2: fractional inputs in other units keep one decimal
                else:
                    c.number_format = "#,##0"
                if f["class"] != "measured":
                    c.fill = KEY
                ws.cell(row=r, column=4, value=f["unit"])
                ws.cell(row=r, column=5, value=f["class"])
                ws.cell(row=r, column=6, value=f["receipt"])
                ws.cell(row=r, column=7, value=f["receipt_text"])
                input_cells[f["id"]] = f"Inputs!$C${r}"
            widths(ws, [8, 34, 12, 8, 11, 9, 70])
        elif sh["kind"] == "options_model":
            opts = sh["options"]
            header(ws, 1, ["Line", "Unit"] + [o["name"] for o in opts] + ["Class", "Receipt"])
            # v1.2: the option rows come from the spec when declared; the option-model default keeps the fixtures byte-identical.
            base = ([(row["label"], row["unit"], row["field"], row.get("format")) for row in sh["rows"]] if sh.get("rows")
                    else [("Capex", "GBPm", "capex", None), ("Annual saving", "GBPm", "annual_saving", None), ("Lead time", "months", "lead_months", None), ("Headcount change", "FTE", "headcount_delta", None)])
            options_layout = {"sheet": sh["name"], "options": [o["id"] for o in opts], "derived_row": len(base) + 3}
            rowmap = {}
            for r, (label, unit, field, fmt) in enumerate(base, start=2):
                ws.cell(row=r, column=1, value=label)
                ws.cell(row=r, column=2, value=unit)
                for c, o in enumerate(opts, start=3):
                    cell = ws.cell(row=r, column=c, value=o[field])
                    cell.font = BLUE
                    cell.fill = KEY
                    cell.number_format = fmt or ("#,##0.0" if unit == "GBPm" else "#,##0")
                ws.cell(row=r, column=3 + len(opts), value="/".join(sorted({o["class"] for o in opts})))
                ws.cell(row=r, column=4 + len(opts), value=", ".join(sorted({o["receipt"] for o in opts})))
                rowmap[field] = r
            r = len(base) + 3
            ws.cell(row=r - 1, column=1, value="Derived").font = BOLD
            for d in sh["derived"]:
                ws.cell(row=r, column=1, value=d["label"])
                for c, o in enumerate(opts, start=3):
                    col = get_column_letter(c)
                    f = d["formula"]
                    for field, rr in rowmap.items():
                        f = f.replace(field, f"{col}${rr}")
                    for fid, ref in input_cells.items():
                        f = f.replace(f"Inputs.{fid}", ref)
                    cell = ws.cell(row=r, column=c, value="=" + f)
                    cell.font = BLACK
                    cell.number_format = d["format"]
                r += 1
            widths(ws, [26, 8] + [30] * len(opts) + [11, 10])
        elif sh["kind"] == "series":
            r = 1
            for s in sh["series"]:
                ws.cell(row=r, column=1, value=f'{s["id"]}  {s["label"]}').font = BOLD
                ws.cell(row=r, column=2, value=f'[{s["class"]}] receipt {s["receipt"]}: {s["receipt_text"]}')
                r += 1
                header(ws, r, [""] + s["categories"])
                for name, vals in s["values"].items():
                    r += 1
                    ws.cell(row=r, column=1, value=name)
                    for c, v in enumerate(vals, start=2):
                        cell = ws.cell(row=r, column=c, value=v)
                        cell.font = BLUE
                        cell.number_format = "0.0%" if s["unit"] == "pct" else "#,##0.00" if s["unit"] == "GBP" else "#,##0.0"
                r += 3
            widths(ws, [28, 22, 22, 22])
        elif sh["kind"] == "checks":
            header(ws, 1, ["Check", "Workbook", "Content model", "Match"])
            checks = []
            if sh.get("checks"):
                # v1.2: declared checks — `formula` over Inputs.<F> cells, or `option` + `derived` row of the options sheet.
                for ch in sh["checks"]:
                    expected, _ = content.calc(ch["calc"], option=ch.get("option"))
                    if "formula" in ch:
                        f = ch["formula"]
                        for fid, ref in input_cells.items():
                            f = f.replace(f"Inputs.{fid}", ref)
                        formula = "=" + f
                    else:
                        col = get_column_letter(3 + options_layout["options"].index(ch["option"]))
                        formula = f"={options_layout['sheet']}!{col}${options_layout['derived_row'] + int(ch.get('derived', 0))}"
                    checks.append((ch["label"], formula, expected))
            else:
                k1, _ = content.calc("K001")
                checks.append(("K001 revenue growth", f"=({input_cells['F003']}-{input_cells['F001']})/{input_cells['F001']}", k1))
                k3, _ = content.calc("K003")
                checks.append(("K003 manual cost premium", f"=({input_cells['F007']}-{input_cells['F006']})/{input_cells['F006']}", k3))
                for i, oid in enumerate(["O1", "O2", "O3"]):
                    v, _ = content.calc("K004", option=oid)
                    col = get_column_letter(3 + i)
                    checks.append((f"K004 payback {oid}", f"=Options!{col}$7", v))
            for r, (label, f, expected) in enumerate(checks, start=2):
                ws.cell(row=r, column=1, value=label)
                ws.cell(row=r, column=2, value=f).number_format = "0.0000"
                ws.cell(row=r, column=3, value=round(expected, 6)).number_format = "0.0000"
                ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<0.0005,"OK","MISMATCH")')
            ws.cell(row=len(checks) + 3, column=1, value="Any MISMATCH is a release blocker (qa.py reads this sheet).")
            widths(ws, [30, 14, 16, 12])

    stamp = dt.datetime.fromisoformat(plan["build_stamp"].replace("Z", "+00:00")).replace(tzinfo=None)
    wb.properties.creator = "doc-studio"
    wb.properties.lastModifiedBy = "doc-studio"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.properties.title = plan["title"]
    wb.save(out)
    return out


if __name__ == "__main__":
    plan = load(sys.argv[1])
    Path(sys.argv[2]).parent.mkdir(parents=True, exist_ok=True)
    render(plan, sys.argv[2], sys.argv[3])
    print(f"xlsx -> {sys.argv[2]}")
