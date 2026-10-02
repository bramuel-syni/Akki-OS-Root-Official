"""Render a resolved paper plan to .docx with python-docx. No design decisions here."""
import io
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from charts import chart_png  # noqa: E402
from studio import load  # noqa: E402


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


class Paper:
    def __init__(self, plan):
        self.p = plan
        self.pal = plan["look"]["palette"]
        self.fonts = plan["look"]["fonts"]
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.2))
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = self.doc.styles["Normal"]
        st.font.name = self.fonts["body"]
        st.font.size = Pt(11)
        st.font.color.rgb = rgb(self.pal["structure"])
        for lvl, size in ((1, 16), (2, 13)):
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = self.fonts["title"]
            h.font.size = Pt(size)
            h.font.bold = lvl == 2
            h.font.color.rgb = rgb(self.pal["structure"])
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), self.fonts["title"])
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")  # python-docx template omits it; validators reject it
        self.footer()

    def footer(self):
        f = self.doc.sections[0].footer.paragraphs[0]
        fu = self.p["furniture"]
        r = f.add_run(f'{fu.get("paper_id", self.p["spec_id"])}   ·   {fu.get("confidentiality", "")}   ·   page ')
        r.font.size = Pt(8)
        r.font.color.rgb = rgb(self.pal["muted"])
        run = f.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(self.pal["muted"])

    def small(self, text, italic=False, mono=False):
        para = self.doc.add_paragraph()
        r = para.add_run(text)
        r.font.size = Pt(8.5)
        r.font.italic = italic
        r.font.color.rgb = rgb(self.pal["muted"])
        if mono:
            r.font.name = self.fonts["number"]
        para.paragraph_format.space_after = Pt(10)
        return para

    def receipt_line(self, cls, ids, text):
        self.small(f"[{cls}]  Receipt {', '.join(ids)}: {text}", mono=True)

    def table(self, header, rows, right_align_from=1, caption=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.font.bold = True
            r.font.size = Pt(9.5)
            shade(c, self.pal["hatch"])
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                r = para.add_run(v)
                r.font.size = Pt(9.5)
                if i >= right_align_from:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    r.font.name = self.fonts["number"]
        if caption:
            self.small(caption, italic=True)
        return t

    def render(self, out):
        d = self.doc
        h = d.add_paragraph()
        r = h.add_run(self.p["title"])
        r.font.name = self.fonts["title"]
        r.font.size = Pt(24)
        r.font.color.rgb = rgb(self.pal["structure"])
        self.small(f'{self.p["subtitle"]}  ·  {self.p["organisation"]}')
        for sec in self.p["sections"]:
            d.add_heading(sec["heading"], level=1)
            for b in sec["blocks"]:
                k = b["kind"]
                if k == "para":
                    d.add_paragraph(b["text"])
                elif k == "figure":
                    png = chart_png(b, self.pal, self.fonts)
                    d.add_picture(io.BytesIO(png), width=Cm(15))
                    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if b.get("caption"):
                        self.small(f'Figure: {b["caption"]}', italic=True)
                    self.receipt_line(b["class"], [b["receipt"]], b["receipt_text"])
                elif k == "table":
                    header = [""] + list(b["values"].keys())
                    rows = [[c] + [f"{b['values'][s][i]:,.2f}" for s in b["values"]] for i, c in enumerate(b["categories"])]
                    self.table(header, rows, caption=b.get("caption"))
                    self.receipt_line(b["class"], [b["receipt"]], b["receipt_text"])
                elif k == "options_table":
                    header = [""] + b["columns"]
                    rows = [[r["label"]] + r["cells"] for r in b["rows"]]
                    self.table(header, rows, caption=b.get("caption"))
                    self.receipt_line(b["class"], b["receipts"], b["receipt_text"])
                elif k == "coverage":
                    d.add_paragraph(f'Covers: {b["covers"]}')
                    self.table(["Cannot say", "Why", "What would close it"],
                               [[g["gap"], g["reason"], g["closes_with"]] for g in b["cannot_say"]], right_align_from=99)
                elif k == "receipts":
                    self.table(["Receipt", "Resolves to"], [[r["id"], r["text"]] for r in b["receipts"]], right_align_from=99)
                elif k == "rows_table":
                    self.small(f'{b["id"]}  {b["label"]}', italic=False)
                    self.table(b["columns"], b["rows"], caption=b.get("caption") or None, right_align_from=b["right_align_from"])
                    if b.get("note"):
                        self.small(b["note"], italic=True)
                    self.receipt_line(b["class"], [b["receipt"]], b["receipt_text"])
                elif k == "glossary":
                    self.table(["Term", "What it means in this pack"], [[t["term"], t["meaning"]] for t in b["terms"]], right_align_from=99)
        cp = d.core_properties
        cp.title = self.p["title"]
        cp.author = "doc-studio"
        cp.last_modified_by = "doc-studio"
        cp.revision = 1
        d.save(out)
        return out


if __name__ == "__main__":
    plan = load(sys.argv[1])
    Path(sys.argv[2]).parent.mkdir(parents=True, exist_ok=True)
    Paper(plan).render(sys.argv[2])
    print(f"docx -> {sys.argv[2]}")
