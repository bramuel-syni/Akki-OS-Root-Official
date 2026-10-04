"""Owner brand catalogue → versioned brand lexicon for the LISTEN-KIK-KAM-SWA commission. The owner's workbook is the authority for what the
listener tracks; every version is recorded with its source hash, what was admitted for matching and what was excluded (and why). The active
version is what the dispatcher ships to the host workers; a change of version re-labels the audio database (admission seals untouched)."""
import hashlib
import re
from collections import OrderedDict

from services.training_lab import store as lab

ALLOWED_SHEETS = {'Product brands': ('family', 0), 'Service brands': ('family', 0), 'Products': ('product_label', 1), 'Services': ('service_provider', 1)}
HEADER_ROW = 5
MIN_ALIAS_CHARS = 3
RULES_VERSION = 2  # alias derivation rules; a rule change on the same workbook mints a new lexicon version
# function words / greetings in English and Kiswahili: a bare alias equal to one of these matches ordinary speech, not a brand (first relabel: "The" 15–22 segments/1000, "jambo" 14.8)
STOPWORDS = {'the', 'and', 'of', 'for', 'with', 'new', 'one', 'a', 'an', 'na', 'ya', 'wa', 'kwa', 'ni', 'tena', 'sana', 'jambo', 'habari', 'leo', 'sasa', 'mimi', 'wewe', 'yeye', 'sisi', 'wao', 'hii', 'hiyo',
             'ile', 'pia', 'bado', 'kama', 'lakini', 'au', 'ama', 'hapa', 'pale', 'basi', 'tu', 'bila', 'kila', 'yote', 'zote', 'wote', 'ndio', 'ndiyo', 'hapana', 'asante', 'karibu', 'pole', 'sawa'}
LEGACY_V1 = [  # hand-written seed list used before the owner catalogue arrived (kept as explicit lineage, flagged legacy_v1 when not in the catalogue)
    {'brand': 'Safaricom', 'aliases': ['safaricom', 'saf']}, {'brand': 'M-Pesa', 'aliases': ['mpesa', 'm-pesa', 'm pesa']}, {'brand': 'Fuliza', 'aliases': ['fuliza']},
    {'brand': 'Airtel', 'aliases': ['airtel', 'airtel money']}, {'brand': 'Equity Bank', 'aliases': ['equity', 'equity bank', 'equitel']}, {'brand': 'KCB', 'aliases': ['kcb', 'kenya commercial bank']},
    {'brand': 'Co-operative Bank', 'aliases': ['co-op bank', 'coop bank', 'cooperative bank', 'co-operative bank']}, {'brand': 'Inooro FM', 'aliases': ['inooro', 'inooro fm']},
    {'brand': 'Kameme FM', 'aliases': ['kameme', 'kameme fm']}, {'brand': 'Coro FM', 'aliases': ['coro fm', 'coro']}, {'brand': 'Citizen', 'aliases': ['citizen', 'citizen tv', 'citizen radio']},
    {'brand': 'Royal Media', 'aliases': ['royal media', 'royal media services']}, {'brand': 'Musyi FM', 'aliases': ['musyi', 'musyi fm']}, {'brand': 'Ramogi FM', 'aliases': ['ramogi']},
    {'brand': 'Radio Jambo', 'aliases': ['radio jambo', 'jambo']}, {'brand': 'Kenya Power', 'aliases': ['kenya power', 'kplc', 'stima']}, {'brand': 'NHIF', 'aliases': ['nhif']},
    {'brand': 'NSSF', 'aliases': ['nssf']}, {'brand': 'KRA', 'aliases': ['kra', 'kenya revenue authority']}, {'brand': 'Huduma', 'aliases': ['huduma', 'huduma centre', 'huduma namba']},
    {'brand': 'Tusker', 'aliases': ['tusker']}, {'brand': 'Coca-Cola', 'aliases': ['coca cola', 'coca-cola', 'coke']}, {'brand': 'Brookside', 'aliases': ['brookside']},
    {'brand': 'Unga', 'aliases': ['jogoo', 'unga limited']}, {'brand': 'Bidco', 'aliases': ['bidco']}, {'brand': 'Betika', 'aliases': ['betika']}, {'brand': 'SportPesa', 'aliases': ['sportpesa', 'sport pesa']},
    {'brand': 'Odibets', 'aliases': ['odibets', 'odi bets']}, {'brand': 'Tala', 'aliases': ['tala']}, {'brand': 'Branch', 'aliases': ['branch loan', 'branch app']}, {'brand': 'KPLC', 'aliases': ['kplc']},
    {'brand': 'Nairobi Water', 'aliases': ['nairobi water']}, {'brand': 'Kenyatta Hospital', 'aliases': ['kenyatta hospital', 'knh']}, {'brand': 'Uber', 'aliases': ['uber']}, {'brand': 'Bolt', 'aliases': ['bolt']},
    {'brand': 'Jumia', 'aliases': ['jumia']}, {'brand': 'KBC', 'aliases': ['kbc']}, {'brand': 'Nation', 'aliases': ['nation media', 'daily nation', 'taifa leo']}, {'brand': 'Standard', 'aliases': ['the standard', 'standard media']},
]


def db():
    database = lab.db()
    database.malpha_brand_lexicon.create_index('version', unique=True)
    return database


def normalise(text):
    return ' '.join(re.sub(r"[^\w\s'’-]", ' ', str(text).lower()).split())


def alias_variants(name):
    """Spoken forms a radio presenter would use for a written label: the label itself, hyphens/apostrophes collapsed, '&' as 'and', trailing legal forms dropped."""
    base = normalise(name)
    legal = re.sub(r'(\s+(plc|ltd|limited|group|kenya|company|holdings|international))+$', '', base)
    legal = re.sub(r'(\s+(of|and|the|ya|wa))+$', '', legal)
    out = OrderedDict()
    for cand in (base, base.replace('-', ' '), base.replace('-', ''), base.replace("'", '').replace('’', ''), re.sub(r'\s*&\s*', ' and ', str(name).lower()), legal):
        cand = normalise(cand)
        if cand: out[cand] = None
    return list(out)


def _matchable(aliases):
    return [a for a in aliases if len(a) >= MIN_ALIAS_CHARS and not a.replace(' ', '').isdigit() and a not in STOPWORDS]


def _stripped(aliases):
    return [a for a in aliases if a in STOPWORDS]


def parse_workbook(path):
    """Workbook → brand entries. Brand families carry sector/company; product labels and service providers are added when not already a family.
    A label that is itself a function word (catalogue normalisation artefact, e.g. label "The" for "The Laughing Cow") takes the offering's leading words."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    entries, seen, sheets = [], {}, {}
    for sheet, (kind, name_col) in ALLOWED_SHEETS.items():
        if sheet not in wb.sheetnames: continue
        rows = [r for i, r in enumerate(wb[sheet].iter_rows(values_only=True), start=1) if i > HEADER_ROW and r and r[name_col]]
        sheets[sheet] = len(rows)
        for r in rows:
            name = str(r[name_col]).strip(); renamed_from = None; offering_alias = None
            if kind != 'family' and normalise(name) in STOPWORDS and r[name_col + 1]:
                words = str(r[name_col + 1]).split()
                if normalise(name) == 'the': renamed_from, name = name, ' '.join(words[:3])  # "The" → "The Laughing Cow"
                elif len(words) >= 2 and normalise(words[1]) not in STOPWORDS: offering_alias = normalise(' '.join(words[:2]))  # "Tena" → alias "tena tissue": the spoken product form
            key = normalise(name)
            if not key: continue
            if key in seen:
                if kind != 'family': seen[key]['offerings'] = seen[key].get('offerings', 0) + 1
                if offering_alias and offering_alias not in seen[key]['aliases']: seen[key]['aliases'].append(offering_alias)
                continue
            entry = {'brand': name, 'kind': kind, 'aliases': alias_variants(name) + ([offering_alias] if offering_alias else [])}
            if renamed_from: entry['renamed_from'] = renamed_from
            if kind == 'family': entry.update(sector=str(r[1] or '').strip() or None, company=str(r[3] or '').strip() or None, source_url=str(r[5] or '').strip() or None)
            else: entry.update(category=str(r[3] or '').strip() or None, offerings=1)
            seen[key] = entry; entries.append(entry)
    for legacy in LEGACY_V1:  # catalogue first; a legacy seed brand not in the catalogue stays visible as legacy_v1 so the owner can strike it
        key = normalise(legacy['brand']); target = seen.get(key) or next((seen[normalise(a)] for a in legacy['aliases'] if normalise(a) in seen), None)
        if target:
            for a in legacy['aliases']:
                if normalise(a) and normalise(a) not in target['aliases']: target['aliases'].append(normalise(a))
        else:
            entry = {'brand': legacy['brand'], 'kind': 'legacy_v1', 'aliases': list(OrderedDict.fromkeys(alias_variants(legacy['brand']) + [normalise(a) for a in legacy['aliases']]))}
            seen[key] = entry; entries.append(entry)
    for e in entries:
        e['matchable_aliases'] = _matchable(e['aliases'])
        e['matchable'] = bool(e['matchable_aliases'])
        stripped = _stripped(e['aliases'])
        if stripped: e['stripped_aliases'] = stripped
        if not e['matchable']:
            e['excluded_reason'] = ('every alias is an everyday word in English/Kiswahili (' + ', '.join(stripped) + ') — would match ordinary speech; needs an owner alias such as "<brand> <product>"') if stripped else \
                f'no alias with ≥{MIN_ALIAS_CHARS} characters that is not a bare number — would match noise in radio speech'
    return entries, sheets


def register(path, filename, actor_id, received_at=None):
    """Record a new lexicon version from the owner workbook; the previous active version is retired (kept). Idempotent on source hash + rules version."""
    data = open(path, 'rb').read(); sha = hashlib.sha256(data).hexdigest()
    existing = db().malpha_brand_lexicon.find_one({'source.sha256': sha, 'rules_version': RULES_VERSION}, {'_id': 0})
    if existing: return existing
    entries, sheets = parse_workbook(path)
    version = (db().malpha_brand_lexicon.find_one(sort=[('version', -1)]) or {}).get('version', 1) + 1
    doc = {'version': version, 'rules_version': RULES_VERSION, 'source': {'filename': filename, 'sha256': sha, 'bytes': len(data), 'sheets': sheets, 'received_at': received_at or lab.now(), 'received_from': actor_id},
           'registered_at': lab.now(), 'active': True, 'brands': entries,
           'counts': {'brands': len(entries), 'matchable': sum(1 for e in entries if e['matchable']), 'excluded': sum(1 for e in entries if not e['matchable']),
                      'by_kind': {k: sum(1 for e in entries if e['kind'] == k) for k in ('family', 'product_label', 'service_provider', 'legacy_v1')}},
           'lexicon_sha256': lexicon_sha(compact(entries))}
    db().malpha_brand_lexicon.update_many({'active': True}, {'$set': {'active': False, 'retired_at': lab.now()}})
    db().malpha_brand_lexicon.insert_one(dict(doc))
    return doc


def compact(entries):
    return [{'brand': e['brand'], 'aliases': e['matchable_aliases']} for e in entries if e.get('matchable')]


def lexicon_sha(compact_entries):
    import json
    return hashlib.sha256(json.dumps(compact_entries, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def active():
    """The lexicon the dispatcher ships. Version 1 is the legacy seed list until the owner catalogue is registered."""
    doc = db().malpha_brand_lexicon.find_one({'active': True}, {'_id': 0})
    if doc: return {'version': doc['version'], 'lexicon': compact(doc['brands']), 'lexicon_sha256': doc['lexicon_sha256'], 'source': doc['source']}
    entries = [{'brand': e['brand'], 'aliases': e['aliases'], 'kind': 'legacy_v1', 'matchable_aliases': _matchable([normalise(a) for a in [e['brand']] + e['aliases']]), 'matchable': True} for e in LEGACY_V1]
    return {'version': 1, 'lexicon': compact(entries), 'lexicon_sha256': lexicon_sha(compact(entries)), 'source': {'filename': None, 'note': 'legacy seed list; no owner catalogue registered yet'}}


def versions():
    return list(db().malpha_brand_lexicon.find({}, {'_id': 0, 'brands': 0}, sort=[('version', 1)]))


def version_detail(version):
    return db().malpha_brand_lexicon.find_one({'version': int(version)}, {'_id': 0})
