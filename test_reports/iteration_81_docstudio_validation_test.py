"""Iteration 81: doc-studio regression + commission/output validation for RMS estate pack.

Scope:
- retyped_numbers selector behavior regression
- producer content/workbook table inclusion + wording checks
- live commission polling (up to 12 minutes) without triggering new producer runs
- conditional publish of the latest clear/released build and static download integrity
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
import requests
from docx import Document
from pptx import Presentation

import sys

VENDORED_ROOT = "/app/backend/vendored/extraction_akki/backend"
if VENDORED_ROOT not in sys.path:
    sys.path.insert(0, VENDORED_ROOT)

from services.transform_forms.doc_studio.producer_estate import build_content, build_workbook  # noqa: E402
from services.transform_forms.doc_studio.publish import publish as publish_released  # noqa: E402
from services.transform_forms.doc_studio.skill.scripts.validate import retyped_numbers  # noqa: E402


COMMISSION_ID = "com-1f3a3606e140"
SOURCE_REF = "eij-1b359b3cbd2e"
DOWNLOAD_BASENAME = "RMS_Estate_and_Language_Review"
BASE_URL = os.environ["REACT_APP_BACKEND_URL"].strip()


def _load_operator_headers() -> dict[str, str]:
    session_path = Path("/tmp/rms_reports_session.json")
    if session_path.exists():
        data = json.loads(session_path.read_text(encoding="utf-8"))
        token = data.get("token")
        if token:
            return {"Authorization": f"Bearer {token}"}

    credentials = re.search(r"\| (demo\.operator@demo\.rms\.example\.com) \| ([^ |]+) \|", Path("/app/memory/test_credentials.md").read_text())
    assert credentials, "Existing operator credentials missing"
    login = requests.post(
        f"{BASE_URL.rstrip('/')}/api/auth/login",
        json={"email": credentials[1], "password": credentials[2]},
        timeout=30,
    )
    assert login.status_code == 200, login.text
    token = login.json().get("access_token")
    assert token, "access_token missing"
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="session")
def auth_headers() -> dict[str, str]:
    return _load_operator_headers()


def test_retyped_numbers_accepts_exact_evidence_columns_only():
    content = {
        "tables": {
            "T040": {
                "columns": [
                    "Store",
                    "Hours held (95 %)",
                    "Speech hours (95 %)",
                ]
            }
        }
    }
    spec_ok = {
        "spec_type": "deck",
        "slides": [
            {
                "type": "table",
                "table_ref": "T040",
                "columns": ["Store", "Hours held (95 %)", "Speech hours (95 %)"]
            }
        ],
    }
    assert retyped_numbers(spec_ok, content) == []

    spec_bad_col = {
        "spec_type": "deck",
        "slides": [
            {"type": "table", "table_ref": "T040", "columns": ["Store", "Hours held (95 %)", "Invented 95 %"]}
        ],
    }
    found = retyped_numbers(spec_bad_col, content)
    assert any("not in evidence table T040" in item for item in found), found

    spec_bad_table = {
        "spec_type": "deck",
        "slides": [
            {"type": "table", "table_ref": "T999", "columns": ["Hours held (95 %)"]}
        ],
    }
    found2 = retyped_numbers(spec_bad_table, content)
    assert any("not in evidence table T999" in item for item in found2), found2


def test_retyped_numbers_keeps_literal_percentage_and_strict_no_content_behavior():
    prose_spec = {
        "spec_type": "paper",
        "sections": [{"heading": "h", "blocks": [{"kind": "para", "para": "Coverage reached 95 % in the latest review."}]}],
    }
    found = retyped_numbers(prose_spec, {"tables": {}})
    assert any("literal number '95'" in item for item in found), found

    # legacy strict mode (content=None) should not treat columns as evidence selectors
    strict_spec = {
        "spec_type": "deck",
        "slides": [{"type": "table", "table_ref": "T040", "columns": ["Hours held (95 %)"]}],
    }
    strict_found = retyped_numbers(strict_spec, None)
    assert any("literal number '95'" in item for item in strict_found), strict_found


def test_producer_content_includes_t045_t046_and_scope_wording():
    report = json.loads(Path("/tmp/rms_estate.json").read_text(encoding="utf-8"))["report"]
    content = build_content(report, "RMS")

    assert "T045" in content["tables"], "T045 Sampling coverage missing"
    assert "T046" in content["tables"], "T046 Readable language estimates missing"

    wb = build_workbook(
        content,
        {"content_ref": "x", "brief_id": "b"},
        {"verdict_id": "v"},
        seed="s",
        stamp="t",
        voice="strategy-house",
    )
    sheet_names = [s["name"] for s in wb["sheets"]]
    assert "Sampling coverage" in sheet_names, sheet_names
    assert "Readable language estimates" in sheet_names, sheet_names

    facts = content.get("facts") or {}
    if "F114" in facts:
        label = (facts["F114"].get("label") or "").lower()
        assert "subset" in label
        assert "not the whole object store" in label

    t013 = content.get("tables", {}).get("T013", {})
    d1_rows = [r for r in (t013.get("rows") or []) if r and r[0] == "D1"]
    assert d1_rows, "D1 decision row missing"
    d1 = d1_rows[0]
    assert "verify cold-bucket storage classes" in str(d1[1]).lower()
    assert "net tiering saving unverified" in str(d1[6]).lower()


@pytest.fixture(scope="session")
def polled_commission(auth_headers):
    """Poll live commission up to 12 min while authoring is in progress."""
    deadline = time.time() + 12 * 60
    history = []
    commission = None

    while time.time() < deadline:
        resp = requests.get(f"{BASE_URL.rstrip('/')}/api/studio/commissions/{COMMISSION_ID}", headers=auth_headers, timeout=45)
        assert resp.status_code == 200, resp.text
        body = resp.json()
        commission = body.get("commission") if isinstance(body, dict) and "commission" in body else body
        release = (commission or {}).get("release") or {}
        state = release.get("state")
        builds = (commission or {}).get("builds") or []
        last = builds[-1] if builds else {}
        history.append(
            {
                "ts": time.time(),
                "state": state,
                "version": commission.get("current_content_version"),
                "last_build_status": last.get("status"),
                "last_build_id": last.get("build_id"),
                "last_build_version": last.get("content_version"),
            }
        )

        if state != "authoring":
            break
        time.sleep(20)

    assert commission is not None
    commission["_poll_history"] = history
    return commission


def test_commission_progress_and_release_quality(polled_commission):
    com = polled_commission
    assert com.get("commission_id") == COMMISSION_ID
    assert com.get("source_ref") == SOURCE_REF

    # Must have authored a newer content version beyond blocked v2.
    assert (com.get("current_content_version") or 0) >= 3, com.get("current_content_version")

    release = com.get("release") or {}
    assert release.get("state") == "clear", {"release": release, "history": com.get("_poll_history")}

    builds = com.get("builds") or []
    assert builds, "no builds on commission"
    last = builds[-1]
    assert last.get("status") == "released", last
    assert (last.get("validate") or {}).get("status") == "PASS", last.get("validate")
    assert (last.get("qa") or {}).get("status") == "PASS", last.get("qa")
    assert (last.get("prove") or {}).get("byte_identical") is True, last.get("prove")

    formats = {o.get("format") for o in (last.get("outputs") or []) if isinstance(o, dict)}
    assert {"docx", "pptx", "xlsx"}.issubset(formats), formats

    # Authoring reviewer blocking disagreements must remain explicit, never silently accepted.
    review = (com.get("authoring") or {}).get("review") or {}
    blocking = review.get("blocking") or []
    resolution = (review.get("resolution") or "").lower()
    silently_accepted = bool(blocking) and ("accepted" in resolution and "without" in resolution and "blocking" in resolution)
    assert not silently_accepted


def test_publish_and_download_hash_validation(polled_commission, auth_headers, tmp_path):
    com = polled_commission
    release = com.get("release") or {}
    builds = com.get("builds") or []
    if release.get("state") != "clear" or not builds:
        pytest.skip("Commission not clear/released; publication must not run.")

    last = builds[-1]
    if last.get("content_version") != com.get("current_content_version"):
        pytest.skip("Last build does not match current content version; refusing publish per requirement.")

    manifest = publish_released("RMS", com)
    assert manifest is not None, "publish() returned None despite clear released build"
    assert manifest.get("build_id") == last.get("build_id")

    for fmt in ("docx", "pptx", "xlsx"):
        entry = (manifest.get("files") or {}).get(fmt)
        assert entry, f"missing manifest entry for {fmt}"
        assert entry["sha256"] == entry["source_sha256"], (fmt, entry)

        public_resp = requests.get(f"{BASE_URL.rstrip('/')}/downloads/{DOWNLOAD_BASENAME}.{fmt}", timeout=60)
        assert public_resp.status_code == 200, (fmt, public_resp.status_code)
        assert hashlib.sha256(public_resp.content).hexdigest() == entry["sha256"], fmt

        local_path = Path("/app/frontend/public/downloads") / f"{DOWNLOAD_BASENAME}.{fmt}"
        assert local_path.exists()
        assert hashlib.sha256(local_path.read_bytes()).hexdigest() == entry["sha256"]

    # Open files with native parsers and sanity-check unresolved tokens/formula errors/scope text.
    docx_bytes = requests.get(f"{BASE_URL.rstrip('/')}/downloads/{DOWNLOAD_BASENAME}.docx", timeout=60).content
    pptx_bytes = requests.get(f"{BASE_URL.rstrip('/')}/downloads/{DOWNLOAD_BASENAME}.pptx", timeout=60).content
    xlsx_bytes = requests.get(f"{BASE_URL.rstrip('/')}/downloads/{DOWNLOAD_BASENAME}.xlsx", timeout=60).content

    doc = Document(io.BytesIO(docx_bytes))
    doc_text = "\n".join(p.text for p in doc.paragraphs if p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if c.text:
                    doc_text += "\n" + c.text

    assert "sampling" in doc_text.lower()
    assert "95 %" in doc_text
    assert ("cannot read" in doc_text.lower()) or ("key refused" in doc_text.lower()) or ("inaccessible" in doc_text.lower())
    assert not re.search(r"\{[FCKOST]\d+", doc_text), "unresolved token found in DOCX"

    pres = Presentation(io.BytesIO(pptx_bytes))
    pptx_text = []
    for slide in pres.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                pptx_text.append(shape.text)
    pptx_joined = "\n".join(pptx_text)
    assert not re.search(r"\{[FCKOST]\d+", pptx_joined), "unresolved token found in PPTX"

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    assert "Sampling coverage" in wb.sheetnames
    assert "Readable language estimates" in wb.sheetnames

    formula_errors = []
    unresolved_tokens = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 2000), values_only=False):
            for cell in row:
                v = cell.value
                if isinstance(v, str):
                    if v.startswith("#") and v in {"#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NUM!"}:
                        formula_errors.append((ws.title, cell.coordinate, v))
                    if re.search(r"\{[FCKOST]\d+", v):
                        unresolved_tokens.append((ws.title, cell.coordinate, v))
    assert not formula_errors, formula_errors[:10]
    assert not unresolved_tokens, unresolved_tokens[:10]

    # Basic OOXML package sanity
    for name, payload in (("docx", docx_bytes), ("pptx", pptx_bytes), ("xlsx", xlsx_bytes)):
        zpath = tmp_path / f"artifact.{name}"
        zpath.write_bytes(payload)
        with zipfile.ZipFile(zpath, "r") as zf:
            assert zf.testzip() is None
